"""Upload validation and plain-text extraction for PDF, DOCX and TXT resumes.

Every check is deterministic and runs before anything is stored:
- the declared extension must match the file's magic bytes (MIME headers are not trusted),
- size, page and character limits bound the work done on untrusted input,
- encrypted PDFs and macro-enabled Word files are rejected.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from typing import Literal

from app.core.errors import AppError

FileKind = Literal["pdf", "docx", "txt"]

MEDIA_TYPES: dict[FileKind, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
}
MAX_PDF_PAGES = 10
MAX_TEXT_CHARS = 60_000
MAX_DOCX_UNCOMPRESSED = 25 * 1024 * 1024  # zip-bomb guard


class UploadRejected(AppError):
    status_code = 422
    code = "UPLOAD_REJECTED"
    message = "This file cannot be used."


class UploadTooLarge(AppError):
    status_code = 413
    code = "UPLOAD_TOO_LARGE"
    message = "The file is too large."


class ExtractionFailed(AppError):
    status_code = 422
    code = "RESUME_EXTRACTION_FAILED"
    message = "We could not read any text from this file. Try a different file or format."


@dataclass(frozen=True)
class Extracted:
    kind: FileKind
    text: str
    pages: int | None
    truncated: bool


def _kind_from_name(filename: str) -> FileKind:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("pdf", "docx", "txt"):
        return ext  # type: ignore[return-value]
    if ext in ("doc", "docm", "dotm", "rtf", "odt", "pages"):
        raise UploadRejected(f"'.{ext}' files are not supported. Upload a PDF, DOCX or TXT file.")
    raise UploadRejected("Unsupported file type. Upload a PDF, DOCX or TXT file.")


def safe_display_name(filename: str) -> str:
    """For display only (storage keys are server-generated)."""
    name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    name = re.sub(r"[^\w.\- ()]+", "_", name).strip(" .")
    return (name or "resume")[:200]


def detect_kind(filename: str, data: bytes) -> FileKind:
    kind = _kind_from_name(filename)
    if kind == "pdf" and not data.startswith(b"%PDF-"):
        raise UploadRejected("This file is not a valid PDF.")
    if kind == "docx":
        if not data.startswith(b"PK\x03\x04"):
            raise UploadRejected("This file is not a valid DOCX document.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
                total = sum(info.file_size for info in archive.infolist())
        except zipfile.BadZipFile:
            raise UploadRejected("This file is not a valid DOCX document.") from None
        if "word/document.xml" not in names:
            raise UploadRejected("This file is not a valid DOCX document.")
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            raise UploadRejected("Macro-enabled documents are not accepted.")
        if total > MAX_DOCX_UNCOMPRESSED:
            raise UploadRejected("This document is too large to process.")
    if kind == "txt" and b"\x00" in data[:4096]:
        raise UploadRejected("This file is not a plain-text file.")
    return kind


def _clean(text: str) -> tuple[str, bool]:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    truncated = len(text) > MAX_TEXT_CHARS
    return text[:MAX_TEXT_CHARS], truncated


def _extract_pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, ValueError, OSError):
        raise UploadRejected("This PDF could not be opened.") from None
    if reader.is_encrypted:
        raise UploadRejected("Password-protected PDFs are not accepted.")
    pages = len(reader.pages)
    if pages > MAX_PDF_PAGES:
        raise UploadRejected(f"Resumes longer than {MAX_PDF_PAGES} pages are not accepted.")
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - a damaged page must not crash the request
            parts.append("")
    return "\n".join(parts), pages


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:  # noqa: BLE001 - python-docx raises many unrelated types
        raise UploadRejected("This DOCX document could not be opened.") from None
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in cells:  # merged cells repeat their text
                    cells.append(value)
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def extract_text(filename: str, data: bytes, max_bytes: int) -> Extracted:
    if len(data) > max_bytes:
        raise UploadTooLarge(f"The file is larger than {max_bytes // (1024 * 1024)} MB.")
    if not data:
        raise UploadRejected("The file is empty.")
    kind = detect_kind(filename, data)
    pages: int | None = None
    if kind == "pdf":
        raw, pages = _extract_pdf(data)
    elif kind == "docx":
        raw = _extract_docx(data)
    else:
        raw = _extract_txt(data)
    text, truncated = _clean(raw)
    if len(re.sub(r"\W", "", text)) < 20:
        raise ExtractionFailed()
    return Extracted(kind=kind, text=text, pages=pages, truncated=truncated)
